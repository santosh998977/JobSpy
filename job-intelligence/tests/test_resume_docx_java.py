import io

from docx import Document

from ai.resume_docx import build_resume_docx


def test_resume_docx_preserves_generated_java_technical_skills():
    docx_bytes = build_resume_docx(
        """
TEST USER
Senior Java Developer
Dallas, TX | email@example.com | 555-555-5555

PROFESSIONAL SUMMARY
Senior Java engineer.

TECHNICAL SKILLS
Languages: Java, SQL
Backend: Spring Boot, Hibernate

PROFESSIONAL EXPERIENCE
Senior Java Developer | Example | Dallas, TX
2020 - Present
- Built Java services.
"""
    )

    document = Document(io.BytesIO(docx_bytes))
    table = next(table for table in document.tables if table.cell(0, 0).text == "Languages")

    assert len(table.rows) == 2
    assert table.cell(0, 1).text == "Java, SQL"
    assert table.cell(1, 0).text == "Backend"
    assert table.cell(1, 1).text == "Spring Boot, Hibernate"
    exported_text = "\n".join(
        cell.text for output_table in document.tables for row in output_table.rows for cell in row.cells
    )
    assert ".NET" not in exported_text
    assert "C#" not in exported_text
